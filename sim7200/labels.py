"""Read-only openpyxl loader over Memory Map.xlsx / Memory Map PLC2.xlsx for GUI labels.

Every sheet in these workbooks that has a "No." header row and an address column
(`Input Address`/`Output Address`/`V Address`) follows the same shape, so this
scans generically instead of hardcoding a sheet list -- new sheets (Stalling,
Stopping Flags, ...) pick up labels automatically. Kept decoupled from
BprDryerSCADA's own catalog/export pipeline (`tools/memory-map-export/`) by
design -- same source files, independent minimal reader.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl

_ADDR_COLUMNS = ("Input Address", "Output Address", "V Address")


@dataclass(frozen=True)
class Label:
    tag: str
    description: str

    @property
    def text(self) -> str:
        return f"{self.tag} - {self.description}" if self.tag else self.description


def _sheet_rows(ws):
    header: tuple | None = None
    for row in ws.iter_rows(values_only=True):
        if header is None:
            if row and row[0] == "No.":
                header = row
            continue
        if not row or row[0] is None:
            continue
        yield dict(zip(header, row))


def load_labels(xlsx_path: str | Path) -> tuple[dict[str, Label], dict[int, Label]]:
    """Returns (address_labels, timer_labels).

    address_labels: normalized operand string ("I0.0", "V1001.0", "VW300", ...) -> Label.
    timer_labels: timer number ("T37" -> 37) -> Label, from the Timers sheet.
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    addr_labels: dict[str, Label] = {}
    timer_labels: dict[int, Label] = {}

    for name in wb.sheetnames:
        for row in _sheet_rows(wb[name]):
            tag = str(row.get("Tag") or row.get("Variable Name") or "")
            desc = str(row.get("Description") or "")
            addr = next((row[c] for c in _ADDR_COLUMNS if row.get(c)), None)
            if addr is not None:
                addr_labels[str(addr).strip().upper()] = Label(tag, desc)
            elif row.get("Timer"):
                timer_labels[int(str(row["Timer"]).strip().lstrip("Tt"))] = Label(tag, desc)
    return addr_labels, timer_labels
